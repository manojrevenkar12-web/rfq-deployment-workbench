"""The workbench on a small simulated world: agents, scoring, discovery, value, deployment, portfolio."""

from __future__ import annotations

import json
from pathlib import Path

from rdw.discovery import analyse_log, compare_live, load_log, log_path, replay, what_if
from rdw.evaluate import similarity, wilson_lower
from rdw.outputs import clarification_letter, compliance_matrix
from rdw.value import value_case


def test_wilson_lower_bound_is_honest_about_small_samples() -> None:
    assert wilson_lower(9, 10) < 0.6
    assert wilson_lower(950, 1000) > 0.93
    assert wilson_lower(0, 0) == 0.0


def test_matching_ignores_numbering_and_case() -> None:
    assert similarity("2.3 Die Spurweite beträgt 1435 mm.", "die spurweite beträgt 1435 mm.") == 1.0
    assert similarity("Die Spurweite beträgt 1435 mm.", "Die Gewährleistung beträgt 24 Monate.") < 0.6


def test_agents_find_most_requirements_and_every_knockout(wb) -> None:  # type: ignore[no-untyped-def]
    for cid in wb.customers():
        m = wb.evaluation(cid)["metrics"]
        assert m["requirement_recall"]["rate"] >= 0.9, cid
        assert m["requirement_precision"]["rate"] >= 0.9, cid
        assert m["unsafe_bid_avoided"]["rate"] == 1.0, cid  # never recommends bidding on a tender that was a no-bid


def test_every_requirement_points_at_its_source(wb) -> None:  # type: ignore[no-untyped-def]
    for a in wb.analyses("gleiswerk"):
        for r in a.requirements:
            assert r.source.file in a.documents and r.source.text and r.source.locator


def test_misses_name_who_fixes_them(wb) -> None:  # type: ignore[no-untyped-def]
    ev = wb.evaluation("brandt")
    owners = {m["owner"] for m in ev["misses"]}
    assert owners <= {"platform", "deployment"}
    # Brandt's profile caps the water tank at 4,000 l; the company built 5,000 l: a deployment fix, not a platform bug.
    tank = [m for m in ev["misses"] if isinstance(m["expected"], dict) and m["expected"]["kind"] == "water_tank_l"]
    assert all(m["owner"] == "deployment" and m["cause"] == "capability_profile" for m in tank)


def test_discovery_finds_where_each_process_really_waits(world: Path) -> None:
    for cid in ("gleiswerk", "voltaro", "brandt"):
        truth = json.loads((world / "truth" / cid / "process.json").read_text())
        found = analyse_log(load_log(log_path(cid, base=world)))
        assert found["bottleneck"]["activity"] == truth["bottleneck"], cid
    brandt = analyse_log(load_log(log_path("brandt", base=world)))
    assert brandt["bottleneck"]["kind"] == "external"


def test_replay_reproduces_the_year_and_agents_only_remove_work(world: Path) -> None:
    cases = load_log(log_path("gleiswerk", base=world))
    same = replay(cases, {})
    observed = {
        c: (evs[-1].complete - evs[0].start).total_seconds() / 86400
        for c, evs in cases.items()
        if evs[-1].activity == "Submitted"
    }
    assert all(abs(same["lead_days"][c] - observed[c]) < 0.01 for c in observed)
    faster = replay(cases, {"Read requirements": 0.3, "Technical review": 0.7})
    assert sum(faster["hours_by_role"].values()) < sum(same["hours_by_role"].values())
    assert all(faster["lead_days"][c] <= same["lead_days"][c] + 1e-6 for c in observed)


def test_buyers_answer_no_faster_with_agents(world: Path) -> None:
    cases = load_log(log_path("brandt", base=world))
    everything = {"Read requirements": 0.0001, "Technical review": 0.0001, "Write documentation": 0.0001}
    floor = replay(cases, everything)
    waits = [
        (evs[i].start - evs[i - 1].complete).total_seconds() / 86400
        for evs in cases.values()
        for i in range(1, len(evs))
        if evs[i].activity == "Buyer answer"
    ]
    assert waits and min(floor["lead_days"].values()) >= 0
    assert sorted(floor["lead_days"].values())[len(floor["lead_days"]) // 2] > 0.5 * sorted(waits)[len(waits) // 2]


def test_value_case_adds_up_and_takes_assumptions(wb) -> None:  # type: ignore[no-untyped-def]
    wi = wb.what_if("voltaro")
    v = value_case(wi)
    assert v["gross_per_year"] == sum(r["value"] for r in v["by_role"])
    assert v["net_per_year"] == v["gross_per_year"] - v["assumptions"]["annual_fee"]
    cheaper = value_case(wi, {"hourly_cost": {"sales_engineering": 10}})
    assert cheaper["gross_per_year"] < v["gross_per_year"]
    nothing = what_if(load_log(log_path("voltaro", base=wb.base)), {}, [])
    assert nothing["hours_saved"] == 0 and value_case(nothing)["by_role"] == []


def test_live_customer_shows_adoption_gaps(wb) -> None:  # type: ignore[no-untyped-def]
    live = wb.live("voltaro")
    assert live is not None
    read = next(r for r in live["activities"] if r["activity"] == "Read requirements")
    assert read["actual_factor"] < 0.6  # the extraction agent took most of the reading
    assert wb.live("brandt") is None


def test_deployment_status_steerco_and_portfolio(wb) -> None:  # type: ignore[no-untyped-def]
    d = wb.deployment("voltaro")
    assert d["health"] in ("green", "amber", "red")
    assert any(m["state"] == "overdue" for m in d["milestones"])
    assert any(s["kind"] == "milestone" for s in d["next_steps"])
    pack = wb.steerco("voltaro")["markdown"]
    for heading in ("## KPIs", "## Since go-live", "## Value", "## Decisions needed", "## Risks"):
        assert heading in pack
    p = wb.portfolio()
    assert {"penalty_cap_percent", "warranty_months", "procedural"} <= set(p["shared_kinds"])
    assert all(c["id"] in ("gleiswerk", "voltaro", "brandt") for c in p["customers"])
    assert p["playbook"]["days_kickoff_to_go_live"] == 49


def test_outputs_for_the_sales_engineer(wb) -> None:  # type: ignore[no-untyped-def]
    from io import BytesIO

    from openpyxl import load_workbook

    a = next(a for a in wb.analyses("voltaro") if any(r.verdict == "clarify" for r in a.requirements))
    sheet = load_workbook(BytesIO(compliance_matrix(a))).active
    assert sheet.max_row == len(a.requirements) + 1
    letter = clarification_letter(a, "Voltaro Ladetechnik GmbH")
    assert ("Bieterfragen" in letter or "Clarification" in letter) and "Voltaro" in letter


def test_reviews_feed_acceptance(wb) -> None:  # type: ignore[no-untyped-def]
    from rdw import store

    a = wb.analyses("voltaro")[0]
    r1, r2 = a.requirements[:2]
    store.save_review("voltaro", a.tender_id, r1.id, r1.kind, r1.verdict, r1.verdict, "", "Jonas")
    store.save_review(
        "voltaro",
        a.tender_id,
        r2.id,
        r2.kind,
        r2.verdict,
        "deviate" if r2.verdict != "deviate" else "comply",
        "no",
        "Jonas",
    )
    store.save_review(
        "voltaro", a.tender_id, r2.id, r2.kind, r2.verdict, r2.verdict, "on second look", "Jonas"
    )  # replaces
    assert store.review_stats("voltaro")["acceptance"] == 1.0
    kpi = next(k for k in wb.kpis("voltaro") if k["key"] == "review_acceptance")
    assert kpi["actual"] == 1.0 and kpi["status"] == "green"


def test_compare_live_lines_up_projection_and_reality(world: Path) -> None:
    rows = compare_live(
        load_log(log_path("voltaro", base=world)),
        load_log(log_path("voltaro", live=True, base=world)),
        {"Read requirements": 0.3},
    )
    read = next(r for r in rows if r["activity"] == "Read requirements")
    assert read["projected_factor"] == 0.3 and abs(read["gap"]) < 0.15
