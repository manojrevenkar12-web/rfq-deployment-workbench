"""Reading a tender as it arrives: PDF specifications, Excel bills of quantities, e-mails, plain text.

Everything becomes a list of segments, one per sentence or table row, each with the file and place it came
from, so every requirement the agents find can be shown in its source.
"""

from __future__ import annotations

import datetime as dt
import email
import email.policy
import re
from dataclasses import dataclass, field
from pathlib import Path

NUMBERED = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+(.*)$")
SENTENCE_END = re.compile(r"(?<=[.!?])\s+(?=[A-ZÄÖÜ„\"(])")
ABBREVIATIONS = (
    "z. B.",
    "z.B.",
    "ggf.",
    "bzw.",
    "inkl.",
    "max.",
    "min.",
    "mind.",
    "ca.",
    "Nr.",
    "e.g.",
    "i.e.",
    "approx.",
)


@dataclass
class Segment:
    doc: str  # spec | lv | eml
    file: str
    locator: str
    text: str
    section: str = ""
    meta: dict[str, str] = field(default_factory=dict)


@dataclass
class Tender:
    id: str
    folder: Path
    segments: list[Segment]
    subject: str = ""
    buyer: str | None = None
    due: str | None = None
    language: str = "de"
    files: list[str] = field(default_factory=list)


def sentences(text: str) -> list[str]:
    protected = text
    for i, a in enumerate(ABBREVIATIONS):
        protected = protected.replace(a, f"\x00{i}\x00")
    parts = SENTENCE_END.split(protected)
    out = []
    for p in parts:
        for i, a in enumerate(ABBREVIATIONS):
            p = p.replace(f"\x00{i}\x00", a)
        p = p.strip()
        if p:
            out.append(p)
    return out


LABEL_LINE = re.compile(r"^[A-ZÄÖÜ][\wÄÖÜäöüß ./-]{2,40}:\s")


def _items(lines: list[tuple[int, str]]) -> list[tuple[int, str, str]]:
    """Joins wrapped lines (also across pages) into items: [(page of first line, number or '', text)]. A line
    continues the previous item unless it is numbered, looks like "Label: value", or the previous item already
    ended a sentence."""
    items: list[list] = []
    for page, raw in lines:
        line = raw.strip()
        if not line:
            continue
        m = NUMBERED.match(line)
        if m and len(m.group(1)) <= 8:
            items.append([page, m.group(1), m.group(2)])
        elif items and not items[-1][2].rstrip().endswith((".", ":", "!", "?")) and not LABEL_LINE.match(line):
            prev = items[-1][2]
            items[-1][2] = prev + line if re.search(r"[A-Za-zÄÖÜäöüß]-$", prev) else prev + " " + line
        else:
            items.append([page, "", line])
    return [(pg, n, re.sub(r"\s+", " ", t)) for pg, n, t in items]


def read_pdf(path: Path) -> tuple[list[Segment], str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    lines = [
        (pno, line) for pno, page in enumerate(reader.pages, 1) for line in (page.extract_text() or "").splitlines()
    ]
    segments: list[Segment] = []
    section = ""
    title = ""
    for pno, number, item in _items(lines):
        if pno == 1 and not title and not number:
            title = item
            continue
        if number and "." not in number and len(item) < 80 and not item.endswith("."):
            section = f"{number} {item}"
            continue
        for s in sentences(item):
            segments.append(Segment("spec", path.name, f"page {pno}" + (f", {number}" if number else ""), s, section))
    return segments, title


def read_text(path: Path) -> tuple[list[Segment], str]:
    lines = path.read_text(encoding="utf-8").splitlines()
    segments: list[Segment] = []
    section, title = "", ""
    for _, number, item in _items([(1, ln.lstrip("#").strip() if ln.startswith("#") else ln) for ln in lines]):
        if not title and not number:
            title = item
            continue
        if number and "." not in number and len(item) < 80 and not item.endswith("."):
            section = f"{number} {item}"
            continue
        for s in sentences(item):
            segments.append(Segment("spec", path.name, number or "text", s, section))
    return segments, title


def read_xlsx(path: Path) -> list[Segment]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    segments = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(c or "").strip().lower() for c in rows[0]]

        def col(*names: str, header: list[str] = header) -> int | None:
            for i, h in enumerate(header):
                if any(n in h for n in names):
                    return i
            return None

        pos, desc, req, cat = (
            col("pos", "item"),
            col("beschreibung", "description"),
            col("anforderung", "requirement"),
            col("kategorie", "category"),
        )
        for row in rows[1:]:
            cells = [str(c).strip() if c is not None else "" for c in row]
            label = cells[desc] if desc is not None else ""
            value = cells[req] if req is not None else ""
            if not (label or value):
                continue
            text = f"{label}: {value}" if label and value else label or value
            meta = {"category": cells[cat]} if cat is not None else {}
            segments.append(
                Segment("lv", path.name, f"row {cells[pos] if pos is not None else '?'}", text, ws.title, meta)
            )
    return segments


def read_eml(path: Path) -> tuple[list[Segment], dict[str, str]]:
    msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    body_part = msg.get_body(preferencelist=("plain",))
    body = body_part.get_content() if body_part else ""
    meta = {"subject": str(msg["Subject"] or ""), "from": str(msg["From"] or ""), "date": str(msg["Date"] or "")}
    segments = []
    for para in re.split(r"\n\s*\n", body):
        flat = " ".join(para.split())
        for s in sentences(flat):
            segments.append(Segment("eml", path.name, "body", s))
    return segments, meta


DUE_DE = re.compile(r"bis zum (\d{1,2})\.(\d{1,2})\.(\d{4})")
DUE_EN = re.compile(r"deadline for offers is (\d{1,2}) (\w+) (\d{4})", re.I)


def _due(segments: list[Segment]) -> str | None:
    for s in segments:
        if s.doc != "eml":
            continue
        if m := DUE_DE.search(s.text):
            return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))).isoformat()
        if m := DUE_EN.search(s.text):
            try:
                return dt.datetime.strptime(" ".join(m.groups()), "%d %B %Y").date().isoformat()
            except ValueError:
                return None
    return None


GERMAN = re.compile(r"\b(der|die|das|und|muss|müssen|ist|sind|mit|für|nach|von|zu)\b", re.I)
ENGLISH = re.compile(r"\b(the|and|shall|must|is|are|with|for|of|to)\b", re.I)


def language(segments: list[Segment]) -> str:
    text = " ".join(s.text for s in segments)
    return "de" if len(GERMAN.findall(text)) >= len(ENGLISH.findall(text)) else "en"


def read_tender(folder: Path) -> Tender:
    segments: list[Segment] = []
    title = ""
    meta: dict[str, str] = {}
    files = []
    for path in sorted(folder.iterdir()):
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            segs, t = read_pdf(path)
            title = title or t
        elif suffix in (".txt", ".md"):
            segs, t = read_text(path)
            title = title or t
        elif suffix == ".xlsx":
            segs = read_xlsx(path)
        elif suffix == ".eml":
            segs, meta = read_eml(path)
        else:
            continue
        files.append(path.name)
        segments.extend(segs)
    subject = title.split(":", 1)[1].strip() if ":" in title else title
    if not subject and meta.get("subject"):
        subject = meta["subject"].split(":", 1)[-1].strip()
    buyer = None
    for s in segments:
        if m := re.match(
            r"(?:Auftraggeber|Contracting authority|Client)\s*:\s*(.+?)(?:\s+(?:Vergabenummer|Reference)\s*:.*)?$",
            s.text,
        ):
            buyer = m.group(1).strip()
            break
    return Tender(folder.name, folder, segments, subject, buyer, _due(segments), language(segments), files)
